"""Проверяем деньги, конкурентные резервы, права и эффекты событий."""
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from io import BytesIO
import uuid
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, func
from polka.api import app
from polka.db import Session
from polka.models import User, Product, Order, OutboxEvent, FulfillmentTask, Notification, SalesFact, Document, StockMovement
from polka.services import create_order
from polka.schemas import Checkout
from polka.events import process_event, publish_one


def create(client, payload, key="checkout-test-001"):
    response = client.post("/api/orders", json=payload, headers={"Idempotency-Key": key})
    assert response.status_code == 201, response.text
    return response.json()


def dispatch(order_id, kind, handler):
    with Session.begin() as db:
        event = db.scalar(select(OutboxEvent).where(OutboxEvent.order_id == order_id, OutboxEvent.kind == kind))
        envelope = {"event_id": event.event_id, "kind": event.kind, **event.payload}
        return process_event(db, handler, envelope)


def test_idempotency_and_decimal_price(clients, checkout_payload):
    first = create(clients["customer"], checkout_payload)
    again = create(clients["customer"], checkout_payload)
    assert first["id"] == again["id"]
    assert first["total"] == "1490.50"
    with Session() as db:
        assert db.get(Product, 1).reserved == 1
        assert db.scalar(select(func.count()).select_from(Order)) == 1
        assert db.scalar(select(func.count()).select_from(OutboxEvent)) == 1


def test_concurrent_last_item(clients, checkout_payload):
    with Session.begin() as db:
        db.get(Product, 1).stock = 1
        ids = db.scalars(select(User.id).where(User.role == "customer")).all()
    barrier = Barrier(2)
    def purchase(user_id):
        with Session() as db:
            user = db.get(User, user_id)
            barrier.wait(timeout=10)
            try:
                create_order(db, user, Checkout(**checkout_payload), str(uuid.uuid4()))
                db.commit()
                return 201
            except HTTPException as error:
                db.rollback()
                return error.status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(purchase, ids))
    assert sorted(results) == [201, 409]
    with Session() as db:
        assert db.get(Product, 1).reserved == 1
        assert db.scalar(select(func.count()).select_from(Order)) == 1


def test_cancel_before_delayed_warehouse_event(clients, checkout_payload):
    order = create(clients["customer"], checkout_payload)
    response = clients["customer"].post(f'/api/orders/{order["id"]}/cancel')
    assert response.status_code == 200
    assert clients["customer"].post(f'/api/orders/{order["id"]}/cancel').status_code == 200
    assert dispatch(order["id"], "order.created", "warehouse") is True
    assert dispatch(order["id"], "order.created", "warehouse") is False
    with Session() as db:
        assert db.get(Product, 1).reserved == 0
        assert db.get(Order, order["id"]).status == "cancelled"
        assert db.scalar(select(func.count()).select_from(FulfillmentTask)) == 0


def test_full_lifecycle_and_duplicate_events(clients, checkout_payload):
    order = create(clients["customer"], checkout_payload)
    oid = order["id"]
    path = f"/api/warehouse/orders/{oid}/status"
    assert clients["warehouse"].post(path, json={"status": "picking"}).status_code == 409
    assert dispatch(oid, "order.created", "warehouse")
    assert not dispatch(oid, "order.created", "warehouse")
    assert clients["warehouse"].post(path, json={"status": "picking"}).status_code == 200
    assert clients["warehouse"].post(path, json={"status": "ready"}).status_code == 409
    assert clients["customer"].post(f"/api/orders/{oid}/cancel").status_code == 409
    item = order["items"][0]
    assert clients["warehouse"].post(f'/api/warehouse/orders/{oid}/items/{item["id"]}', json={"picked": 1}).status_code == 200
    for status in ("ready", "shipped", "shipped"):
        response = clients["warehouse"].post(path, json={"status": status})
        assert response.status_code == 200, response.text
    assert clients["warehouse"].post(path, json={"status": "delivered"}).status_code == 403
    assert clients["manager"].post(path, json={"status": "delivered"}).status_code == 200
    for handler in ("notifications", "analytics", "documents"):
        assert dispatch(oid, "order.delivered", handler)
        assert not dispatch(oid, "order.delivered", handler)
    with Session() as db:
        assert (db.get(Product, 1).stock, db.get(Product, 1).reserved) == (2, 0)
        for model in (FulfillmentTask, Notification, SalesFact, StockMovement):
            assert db.scalar(select(func.count()).select_from(model)) == 1
        assert db.get(Document, oid).content.startswith(b"%PDF")
    pdf = clients["customer"].get(f"/api/orders/{oid}/document.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    pdf_head = clients["customer"].head(f"/api/orders/{oid}/document.pdf")
    assert pdf_head.status_code == 200 and not pdf_head.content
    excel = clients["manager"].get("/api/manager/reports.xlsx")
    assert excel.status_code == 200
    excel_head = clients["manager"].head("/api/manager/reports.xlsx")
    assert excel_head.status_code == 200 and not excel_head.content
    sheet = load_workbook(BytesIO(excel.content)).active
    assert sheet.max_row == 2 and sheet["E2"].value == 1490.5
    assert clients["manager"].get("/api/manager/dashboard").json()["revenue"] == "1490.50"


def test_permissions_csrf_and_untrusted_fields(clients, checkout_payload):
    order = create(clients["customer"], checkout_payload)
    assert clients["other"].get(f'/api/orders/{order["id"]}').status_code == 404
    assert clients["other"].post(f'/api/orders/{order["id"]}/cancel').status_code == 404
    assert clients["customer"].get("/api/warehouse/orders").status_code == 403
    assert clients["warehouse"].get("/api/manager/products").status_code == 403
    assert clients["warehouse"].post("/api/orders", json=checkout_payload, headers={"Idempotency-Key": "staff-invalid"}).status_code == 403
    with TestClient(app) as guest:
        assert guest.get("/api/orders").status_code == 401
        response = guest.post("/api/auth/register", json={"email": "new@test.local", "name": "Новый покупатель", "password": "password1234", "role": "manager"})
        assert response.status_code == 422
    token = clients["customer"].headers.pop("X-CSRF-Token")
    assert clients["customer"].post(f'/api/orders/{order["id"]}/cancel').status_code == 403
    clients["customer"].headers["X-CSRF-Token"] = token
    assert clients["customer"].post(f'/api/orders/{order["id"]}/cancel', headers={"Origin": "https://foreign.example"}).status_code == 403
    manipulated = {**checkout_payload, "total": "1.00"}
    assert clients["customer"].post("/api/orders", json=manipulated, headers={"Idempotency-Key": "price-fraud-test"}).status_code == 422


def test_stock_cannot_drop_below_reserved_and_archive_keeps_order(clients, checkout_payload):
    order = create(clients["customer"], checkout_payload)
    product = clients["manager"].get("/api/manager/products").json()[0]
    data = {key: product[key] for key in ("sku", "name", "description", "category", "image", "price", "stock", "active")}
    assert clients["manager"].put("/api/manager/products/1", json={**data, "stock": 0}).status_code == 409
    assert clients["manager"].put("/api/manager/products/1", json={**data, "active": False}).status_code == 200
    assert clients["customer"].get("/api/products").json()["products"] == []
    assert clients["customer"].get(f'/api/orders/{order["id"]}').json()["items"][0]["name"] == "Тестовая ваза"


def test_failed_publication_keeps_event_and_order_sequence(clients, checkout_payload):
    order = create(clients["customer"], checkout_payload)
    clients["customer"].post(f'/api/orders/{order["id"]}/cancel')
    class UnavailableBroker:
        def produce(self, *args, **kwargs):
            raise BufferError("Брокер недоступен")
    assert publish_one(UnavailableBroker())
    with Session() as db:
        first = db.scalar(select(OutboxEvent).order_by(OutboxEvent.id))
        assert first.published_at is None and first.attempts == 1
        assert db.get(Order, order["id"]).status == "cancelled"
        assert db.get(Product, 1).reserved == 0
    # Следующий статус того же заказа не обгоняет неотправленное событие.
    assert publish_one(UnavailableBroker()) is False
